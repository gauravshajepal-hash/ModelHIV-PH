from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from epigraph_ph.runtime import RunContext, ensure_dir, sha256_file, utc_now_iso, write_json

_DEFAULT_WORKBOOK_PATH = Path(__file__).resolve().parent / "P_Data_Extract_From_World_Development_Indicators.xlsx"
_HIV_SERIES_PREFIXES = ("SH.HIV", "SH.DYN.AIDS")


def default_wdi_hiv_workbook_path() -> Path:
    return _DEFAULT_WORKBOOK_PATH


def _is_hiv_series(series_code: str, series_name: str) -> bool:
    code = str(series_code or "").strip().upper()
    name = str(series_name or "").strip()
    return code.startswith(_HIV_SERIES_PREFIXES) or bool(re.search(r"\b(HIV|AIDS)\b", name, flags=re.IGNORECASE))


def _parse_year_header(value: Any) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"(\d{4})", text)
    if not match:
        return None
    return int(match.group(1))


def _safe_metric_slug(series_code: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "_", str(series_code or "").strip().lower()).strip("_")
    return f"wdi_{cleaned}" if cleaned else "wdi_unknown_series"


def _normalize_unit(series_name: str, metadata_unit: str | None) -> str:
    unit_text = str(metadata_unit or "").strip().lower()
    name_text = str(series_name or "").strip().lower()
    if "%" in name_text or unit_text in {"percent", "%"}:
        return "percent"
    if "per 1,000" in name_text:
        return "rate_per_1000"
    if unit_text == "number":
        return "count_people"
    return "scalar"


def _series_measurement_class(series_code: str, series_name: str) -> str:
    if str(series_code or "").startswith("SH."):
        return "model_estimate"
    if "hiv" in str(series_name or "").lower() or "aids" in str(series_name or "").lower():
        return "model_estimate"
    return "external_series"


def _coerce_numeric(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _metadata_by_code(workbook_path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Series - Metadata" not in workbook.sheetnames:
        return {}
    sheet = workbook["Series - Metadata"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows)]
    metadata: dict[str, dict[str, Any]] = {}
    for row in rows:
        payload = {
            str(column): row[idx]
            for idx, column in enumerate(header)
            if idx < len(row) and str(column).strip()
        }
        code = str(payload.get("Code") or "").strip()
        if code:
            metadata[code] = payload
    return metadata


def extract_wdi_hiv_rows(
    workbook_path: str | Path,
    *,
    country_code: str = "PHL",
    country_name: str = "Philippines",
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook_file = Path(workbook_path)
    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    data_sheet = workbook["Data"]
    metadata_map = _metadata_by_code(workbook_file)

    row_iter = data_sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(row_iter)]
    year_columns: list[tuple[int, int]] = []
    for index, value in enumerate(header):
        year = _parse_year_header(value)
        if year is not None:
            year_columns.append((index, year))

    metric_rows: list[dict[str, Any]] = []
    series_summary_rows: list[dict[str, Any]] = []

    for row in row_iter:
        if not row or len(row) < 4:
            continue
        series_name = str(row[0] or "").strip()
        series_code = str(row[1] or "").strip()
        row_country_name = str(row[2] or "").strip()
        row_country_code = str(row[3] or "").strip()
        if row_country_code != country_code:
            continue
        if not _is_hiv_series(series_code, series_name):
            continue

        metadata = metadata_map.get(series_code, {})
        unit = _normalize_unit(series_name, str(metadata.get("Unit of measure") or ""))
        measurement_class = _series_measurement_class(series_code, series_name)
        values_for_series: list[tuple[int, float]] = []
        metric_name = _safe_metric_slug(series_code)

        for index, year in year_columns:
            if index >= len(row):
                continue
            numeric_value = _coerce_numeric(row[index])
            if numeric_value is None:
                continue
            values_for_series.append((year, numeric_value))
            metric_rows.append(
                {
                    "year": year,
                    "time": f"{year:04d}-01",
                    "metric_name": metric_name,
                    "value": numeric_value,
                    "reported_value": str(row[index]),
                    "unit": unit,
                    "measurement_class": measurement_class,
                    "series_kind": "annual_wdi_series",
                    "temporal_precision": "annual_series",
                    "geo": row_country_name or country_name,
                    "region": "national",
                    "province": row_country_name or country_name,
                    "source_id": f"wdi_{series_code.lower().replace('.', '_')}",
                    "source_label": "World Development Indicators HIV Series",
                    "source_url": "https://databank.worldbank.org/source/world-development-indicators",
                    "source_note": str(metadata.get("Source") or ""),
                    "source_dataset": str(metadata.get("Dataset") or "WB_WDI"),
                    "source_organization": str(metadata.get("Source") or ""),
                    "series_code": series_code,
                    "series_name": series_name,
                    "evidence_confidence": 0.97,
                    "value_semantics": "point_estimate",
                }
            )

        if values_for_series:
            years = [year for year, _ in values_for_series]
            series_summary_rows.append(
                {
                    "series_code": series_code,
                    "series_name": series_name,
                    "metric_name": metric_name,
                    "country_name": row_country_name or country_name,
                    "country_code": row_country_code,
                    "unit": unit,
                    "measurement_class": measurement_class,
                    "year_count": len(values_for_series),
                    "min_year": min(years),
                    "max_year": max(years),
                    "source_dataset": str(metadata.get("Dataset") or "WB_WDI"),
                    "source_note": str(metadata.get("Source") or ""),
                    "periodicity": str(metadata.get("Periodicity") or ""),
                }
            )

    metric_rows.sort(key=lambda item: (str(item["metric_name"]), int(item["year"])))
    series_summary_rows.sort(key=lambda item: str(item["series_code"]))
    return metric_rows, series_summary_rows


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    if not rows:
        with path.open("w", encoding="utf-8", newline="") as handle:
            handle.write("")
        return
    fieldnames = sorted({key for row in rows for key in row.keys()})
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def run_harp_archive_wdi_hiv_extract(
    *,
    run_id: str,
    plugin_id: str = "hiv",
    workbook_path: str | Path | None = None,
    country_code: str = "PHL",
    country_name: str = "Philippines",
) -> dict[str, Any]:
    workbook_file = Path(workbook_path) if workbook_path else default_wdi_hiv_workbook_path()
    if not workbook_file.exists():
        raise FileNotFoundError(f"WDI workbook not found: {workbook_file}")

    ctx = RunContext.create(run_id=run_id, plugin_id=plugin_id)
    artifact_dir = ensure_dir(Path(ctx.run_dir) / "harp_archive_wdi")

    metric_rows, series_rows = extract_wdi_hiv_rows(
        workbook_file,
        country_code=country_code,
        country_name=country_name,
    )
    metrics_json_path = artifact_dir / "wdi_hiv_metric_rows.json"
    metrics_csv_path = artifact_dir / "wdi_hiv_metric_rows.csv"
    inventory_json_path = artifact_dir / "wdi_hiv_series_inventory.json"
    inventory_csv_path = artifact_dir / "wdi_hiv_series_inventory.csv"

    write_json(metrics_json_path, metric_rows)
    write_json(inventory_json_path, series_rows)
    _write_csv(metrics_csv_path, metric_rows)
    _write_csv(inventory_csv_path, series_rows)

    manifest = {
        "run_id": run_id,
        "plugin_id": plugin_id,
        "generated_at": utc_now_iso(),
        "workbook_path": str(workbook_file),
        "workbook_sha256": sha256_file(workbook_file),
        "country_code": country_code,
        "country_name": country_name,
        "series_count": len(series_rows),
        "metric_row_count": len(metric_rows),
        "artifact_paths": {
            "metric_rows_json": str(metrics_json_path),
            "metric_rows_csv": str(metrics_csv_path),
            "series_inventory_json": str(inventory_json_path),
            "series_inventory_csv": str(inventory_csv_path),
        },
    }
    write_json(artifact_dir / "wdi_hiv_manifest.json", manifest)
    return manifest
